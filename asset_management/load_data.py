"""读取股票、客户、新闻和经济指标数据，不修改输入文件。

将工作簿布局和源字段转换为分析模块共用的表格接口，并在入口检查
日期、资产标签及数值类型。价格和指标允许缺失值留待特征层处理；
客户分群要求完整数值记录，因此客户缺失值在读取阶段直接报错。
所有函数仅访问显式指定的本地路径，不自动下载或补写原始数据。
"""

from __future__ import annotations

import json
from pathlib import Path

import numpy as np
import pandas as pd
from openpyxl import load_workbook


def _sheet_rows(path: Path, preferred: str | None, fallback: int) -> list[tuple]:
    """读取首选工作表，不存在时按从零开始的备用位置读取。

    返回包含表头在内的单元格值列表。只读模式避免修改工作簿；
    ``data_only`` 读取文件已保存的公式计算值，不在此重新计算公式。
    即使读取出错也关闭工作簿，使后续运行或 Excel 使用不受句柄占用。
    """
    if not path.is_file():
        raise FileNotFoundError(f"Missing input file: {path}")
    try:
        workbook = load_workbook(path, read_only=True, data_only=True)
        try:
            sheet = workbook[preferred] if preferred in workbook.sheetnames else workbook.worksheets[fallback]
            return list(sheet.iter_rows(values_only=True))
        finally:
            workbook.close()
    except (OSError, ValueError, IndexError, KeyError) as exc:
        raise ValueError(f"Cannot read the expected worksheet in {path.name}: {exc}") from exc


def _numeric(frame: pd.DataFrame, context: str) -> pd.DataFrame:
    """统一已知缺失标记与数值类型，保留索引和列名。

    空字符串、短横线和两类不可用标记转换为缺失值；其他非数值文本
    明确报错，避免将字段错位或格式问题静默当作缺失观测。返回浮点
    表格，不填补缺失值；无限值无法用于后续统计，因此直接拒绝。
    """
    missing_tokens = {"": np.nan, "-": np.nan, "N/A": np.nan, "#N/A": np.nan}
    result = frame.replace(missing_tokens)
    try:
        result = result.apply(pd.to_numeric, errors="raise").astype(float)
    except (ValueError, TypeError) as exc:
        raise ValueError(f"{context} contains a nonnumeric value outside the supported missing-value markers.") from exc
    if np.isinf(result.to_numpy()).any():
        raise ValueError(f"{context} contains infinite values.")
    return result


def _date_index(frame: pd.DataFrame, context: str) -> pd.DataFrame:
    """将索引转为名为 date 的日期索引，验证后按时间升序返回。

    日期不得为空或重复：重复观测会使相邻收益含义不明确，也可能让
    月度合并重复计入同一期。此处修改传入表格的索引，不补齐缺失日期。
    """
    try:
        frame.index = pd.DatetimeIndex(pd.to_datetime(frame.index, errors="raise"), name="date")
    except (ValueError, TypeError) as exc:
        raise ValueError(f"{context} contains an invalid date.") from exc
    if frame.index.isna().any() or frame.index.has_duplicates:
        raise ValueError(f"{context} requires nonmissing, unique dates.")
    return frame.sort_index()


def load_prices(path: str | Path) -> pd.DataFrame:
    """根据彭博数据的两行表头提取股票 PX_LAST 价格列。

    第一行包含资产名称，第二行包含字段名称；合并表头后的空白位置
    沿用当前资产分组，只选择 Equity 分组中的 PX_LAST 收盘价字段。
    资产代码取名称首段，因此市场指数及同一资产的其他数据字段不会
    混入价格矩阵。返回日期升序、股票代码为列名的浮点表格。

    删除完全空白的行，但保留个别价格缺失值，避免收益率计算跨越
    缺失观测。每只股票至少须有两个观测价格；是否存在共同有效的
    相邻日期由收益率函数进一步验证。已观测价格须为正数。
    """
    path = Path(path)
    rows = _sheet_rows(path, "Bloomberg raw", 1)
    if len(rows) < 3:
        raise ValueError(f"{path.name} must contain two header rows and price observations.")
    selected: list[tuple[int, str]] = []
    group = ""
    # 逐列继承资产分组，而不是假定每只股票在工作簿中占固定列数。
    for position, (label, field) in enumerate(zip(rows[0], rows[1])):
        if label is not None:
            group = str(label).strip()
        if "Equity" in group and str(field).strip() == "PX_LAST":
            selected.append((position, group.split()[0]))
    if not selected or len({name for _, name in selected}) != len(selected):
        raise ValueError(f"{path.name} must provide one PX_LAST column per unique equity ticker.")
    observations = [row for row in rows[2:] if any(value is not None for value in row)]
    result = pd.DataFrame(
        [[row[position] for position, _ in selected] for row in observations],
        index=[row[0] for row in observations],
        columns=[name for _, name in selected],
    )
    result = _date_index(_numeric(result, "Stock prices"), "Stock prices")
    if result.empty or result.notna().sum().min() < 2:
        raise ValueError("Each stock needs at least two observed prices.")
    if (result <= 0).any().any():
        raise ValueError("Observed stock prices must be strictly positive.")
    return result


def load_clients(path: str | Path) -> pd.DataFrame:
    """读取客户数值特征，以唯一 client_ID 作为行索引。

    从 Data 工作表或备用第二张表读取首行字段名，移除完全空白的
    行列。必须包含客户编号、risk_profile 和 age_group；股票持仓
    列中的完整 Equity 标签转换为代码，以便与其他模块统一命名。
    返回的其余字段全部为浮点数，不自行插补、删除不完整客户或缩放
    特征；出现缺失记录时要求先明确处理，避免改变客户分群的含义。
    """
    path = Path(path)
    rows = _sheet_rows(path, "Data", 1)
    if len(rows) < 2:
        raise ValueError(f"{path.name} does not contain client records.")
    frame = pd.DataFrame(rows[1:], columns=rows[0]).dropna(axis=1, how="all").dropna(axis=0, how="all")
    required = {"client_ID", "risk_profile", "age_group"}
    if not required.issubset(frame.columns):
        raise ValueError(f"{path.name} requires columns: {', '.join(sorted(required))}.")
    frame = frame.set_index("client_ID")
    if frame.index.isna().any() or frame.index.has_duplicates:
        raise ValueError("Client IDs must be present and unique.")
    frame.columns = [str(column).split()[0] if "Equity" in str(column) else str(column) for column in frame.columns]
    # 标签简化可能把不同原始字段变成同一代码，须在计算前发现冲突。
    if frame.columns.has_duplicates:
        raise ValueError("Client features contain duplicate column names after ticker normalization.")
    frame = _numeric(frame, "Client features")
    if frame.empty or frame.isna().any().any():
        raise ValueError("Client features require complete numeric records; fill or remove incomplete rows explicitly.")
    return frame


def load_news(path: str | Path) -> pd.DataFrame:
    """读取新闻 JSON 表，按输入顺序返回四个固定字段。

    字段依次为 Equity、Source、Date/Time、Headline；清除原行索引后
    使用连续行号，股票标签简化为代码。编码支持 UTF-8 的可选字节序
    标记，股票代码和标题必须是非空字符串。

    日期字段保留源文本，不解析、排序或去重。情绪模块将每只股票的
    前若干条当作近期窗口，因此文件的排列顺序会直接影响调整结果。
    """
    path = Path(path)
    if not path.is_file():
        raise FileNotFoundError(f"Missing input file: {path}")
    try:
        with path.open(encoding="utf-8-sig") as handle:
            frame = pd.DataFrame(json.load(handle))
    except (ValueError, TypeError) as exc:
        raise ValueError(f"{path.name} must be a JSON table of headline records.") from exc
    required = ["Equity", "Source", "Date/Time", "Headline"]
    if not set(required).issubset(frame.columns):
        raise ValueError(f"{path.name} requires columns: {', '.join(required)}.")
    frame = frame.loc[:, required].copy().reset_index(drop=True)
    if frame.empty or frame[["Equity", "Headline"]].isna().any().any():
        raise ValueError("News must contain nonmissing equity tickers and headlines.")
    for column in ("Equity", "Headline"):
        if not frame[column].map(lambda value: isinstance(value, str) and bool(value.strip())).all():
            raise ValueError(f"News {column} values must be nonempty strings.")
    frame["Equity"] = frame["Equity"].str.strip().str.split().str[0]
    return frame


def load_indicators(path: str | Path) -> pd.DataFrame:
    """提取月度指标数据，排除分类行和季度指标。

    工作表中 Monthly Indicators 行给出横向日期，下方各行给出不同
    指标。读取时转置为日期索引、指标名称为列的浮点表格，跳过无数值
    的分类标题；遇到下一个指标分区立即停止，不混入季度指标。

    将交易月末日期统一为自然月末，便于与月度股票收益按同月对齐；
    保留缺失值，不进行跨期填补。统一月末后若同月存在重复记录则报错，
    防止一个月在回归中被重复使用。
    """
    path = Path(path)
    rows = _sheet_rows(path, None, 0)
    starts = [i for i, row in enumerate(rows) if str(row[0]).strip().lower() == "monthly indicators"]
    if len(starts) != 1:
        raise ValueError(f"{path.name} must contain exactly one 'Monthly Indicators' section.")
    start = starts[0]
    header = rows[start]
    positions = [i for i, value in enumerate(header[1:], 1) if value is not None]
    try:
        dates = pd.to_datetime([header[i] for i in positions], errors="raise")
    except (ValueError, TypeError) as exc:
        raise ValueError(f"{path.name} has invalid monthly date headers.") from exc
    values: dict[str, list] = {}
    # 空分类行不构成观测；下一个频率分区是当前月度区域的结束标志。
    for row in rows[start + 1 :]:
        label = str(row[0]).strip() if row[0] is not None else ""
        if label.lower().endswith(" indicators"):
            break
        series = [row[i] for i in positions]
        if not label or all(value is None for value in series):
            continue
        if label in values:
            raise ValueError(f"Duplicate monthly indicator: {label}")
        values[label] = series
    if not values or not positions:
        raise ValueError(f"{path.name} contains no monthly indicator observations.")
    frame = _numeric(pd.DataFrame(values, index=dates), "Monthly indicators")
    # 使用所属月份定位自然月末，周末不会把本月指标错误推到下个月。
    frame.index = frame.index.to_period("M").to_timestamp("M")
    return _date_index(frame, "Monthly indicators")


def load_inputs(raw_dir: str | Path) -> dict[str, pd.DataFrame]:
    """按约定文件名读取完整模式的四类本地输入。

    返回字典中的 prices 为日期 × 股票价格，clients 为客户 × 数值
    特征，news 为逐条新闻记录，indicators 为月份 × 经济指标。
    各读取函数只负责本类数据校验，跨类股票匹配及日期交集由后续模型
    完成。任一文件缺失或格式不合法时立即抛出错误，不返回部分结果。
    """
    raw_dir = Path(raw_dir)
    return {
        "prices": load_prices(raw_dir / "ASX200top10.xlsx"),
        "clients": load_clients(raw_dir / "Client_Details.xlsx"),
        "news": load_news(raw_dir / "news_dump.json"),
        "indicators": load_indicators(raw_dir / "Economic_Indicators.xlsx"),
    }


def load_sample_prices(sample_dir: str | Path) -> pd.DataFrame:
    """读取演示用 prices.csv，返回日期升序的资产价格矩阵。

    文件必须包含 date 和至少一个资产价格字段；日期转为索引，其余
    字段转为浮点数。此函数验证已观测价格为正，不补齐缺失日期或价格；
    后续收益率步骤继续检查是否存在所有资产共同有效的相邻观测。
    示例数据与完整模式的工作簿入口分开，便于演示读取和报告生成。
    """
    path = Path(sample_dir) / "prices.csv"
    if not path.is_file():
        raise FileNotFoundError(f"Missing synthetic sample prices: {path}")
    frame = pd.read_csv(path)
    if "date" not in frame.columns or len(frame.columns) < 2:
        raise ValueError("Sample prices.csv requires a date column and at least one asset column.")
    frame = _date_index(_numeric(frame.set_index("date"), "Sample prices"), "Sample prices")
    if frame.empty or (frame <= 0).any().any():
        raise ValueError("Sample prices must contain strictly positive observed prices.")
    return frame
